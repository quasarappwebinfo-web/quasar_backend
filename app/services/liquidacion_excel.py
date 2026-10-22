"""Excel de liquidación: días con asistencia × valor día, y el detalle día a día."""
from __future__ import annotations

from collections import defaultdict
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image as PilImage
from sqlalchemy.orm import Session, joinedload

from app.models import Obra, ObraDailyAttendance, Worker

NAVY = "111111"
BLUE = "014CBD"
ZEBRA = "F7F9FC"
LINE = "D5DDE8"
WHITE = "FFFFFF"
MUTED = "5C6670"
GOLD = "C4A35A"

MESES = (
    "enero",
    "febrero",
    "marzo",
    "abril",
    "mayo",
    "junio",
    "julio",
    "agosto",
    "septiembre",
    "octubre",
    "noviembre",
    "diciembre",
)

_LOGO = Path(__file__).resolve().parents[1] / "assets" / "quasar-logo.png"
_THIN = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)
_MONEY = '"$"#,##0'


def _fecha(value: date) -> str:
    return f"{value.day} de {MESES[value.month - 1]} de {value.year}"


def _logo_png() -> Optional[str]:
    """El PNG trae fondo negro. Se vuelve transparente para sentarlo en la franja oscura."""
    if not _LOGO.is_file():
        return None
    image = PilImage.open(_LOGO).convert("RGBA")
    pixels = []
    for red, green, blue, alpha in image.getdata():
        if red < 28 and green < 28 and blue < 28:
            pixels.append((red, green, blue, 0))
        else:
            pixels.append((red, green, blue, alpha))
    image.putdata(pixels)
    box = image.getbbox()
    if box:
        image = image.crop(box)
    target = Path("/tmp/quasar-logo-excel.png")
    image.save(target, format="PNG")
    return str(target)


def _paint(ws, row: int, start: int, end: int, fill: PatternFill) -> None:
    for col in range(start, end + 1):
        ws.cell(row, col).fill = fill


def _header_band(ws, last_col: int, title: str, subtitle: str) -> None:
    dark = PatternFill("solid", fgColor=NAVY)
    for row in range(1, 6):
        ws.row_dimensions[row].height = 18
        _paint(ws, row, 1, last_col, dark)
    ws.row_dimensions[2].height = 26
    ws.merge_cells(start_row=2, start_column=3, end_row=2, end_column=last_col)
    ws.merge_cells(start_row=3, start_column=3, end_row=3, end_column=last_col)
    ws.merge_cells(start_row=4, start_column=3, end_row=4, end_column=last_col)
    title_cell = ws.cell(2, 3, title)
    title_cell.font = Font(name="Calibri", size=20, bold=True, color=WHITE)
    title_cell.alignment = Alignment(vertical="center")
    brand = ws.cell(3, 3, "QUASAR  ·  Arquitectura, diseño y construcción")
    brand.font = Font(name="Calibri", size=11, color=GOLD)
    period = ws.cell(4, 3, subtitle)
    period.font = Font(name="Calibri", size=11, color="D6D6D6")
    logo = _logo_png()
    if logo is not None:
        picture = XlImage(logo)
        picture.height = 72
        picture.width = 118
        picture.anchor = "A1"
        ws.add_image(picture)


def _table_header(ws, row: int, labels: List[str]) -> None:
    fill = PatternFill("solid", fgColor=BLUE)
    font = Font(name="Calibri", size=11, bold=True, color=WHITE)
    for index, label in enumerate(labels, start=1):
        cell = ws.cell(row, index, label)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _THIN
    ws.row_dimensions[row].height = 28
    ws.freeze_panes = f"A{row + 1}"
    ws.print_title_rows = f"{row}:{row}"


def _write_row(ws, row: int, values: list, *, zebra: bool, money_cols: Tuple[int, ...] = ()) -> None:
    fill = PatternFill("solid", fgColor=ZEBRA if zebra else WHITE)
    font = Font(name="Calibri", size=11, color="1A1C1F")
    for index, value in enumerate(values, start=1):
        cell = ws.cell(row, index, value)
        cell.fill = fill
        cell.font = font
        cell.border = _THIN
        cell.alignment = Alignment(vertical="center", wrap_text=index in (2, 7))
        if index in money_cols and isinstance(value, (int, float, Decimal)):
            cell.number_format = _MONEY
            cell.alignment = Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[row].height = 22


def _totals_row(ws, row: int, last_col: int, label: str, days, money) -> None:
    fill = PatternFill("solid", fgColor="10233F")
    font = Font(name="Calibri", size=11, bold=True, color=WHITE)
    for col in range(1, last_col + 1):
        cell = ws.cell(row, col, None)
        cell.fill = fill
        cell.font = font
        cell.border = _THIN
    ws.cell(row, 1, label).font = font
    days_cell = ws.cell(row, 6, days)
    days_cell.font = font
    days_cell.alignment = Alignment(horizontal="center")
    money_cell = ws.cell(row, 8, money)
    money_cell.font = font
    money_cell.number_format = _MONEY
    money_cell.alignment = Alignment(horizontal="right")
    ws.row_dimensions[row].height = 24


def _widths(ws, widths: List[float]) -> None:
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _sheet_setup(ws, landscape: bool) -> None:
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_TABLOID if landscape else ws.PAPERSIZE_LETTER
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.horizontalCentered = True
    ws.page_margins.left = 0.5
    ws.page_margins.right = 0.5
    ws.page_margins.top = 0.6
    ws.page_margins.bottom = 0.6
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 110
    ws.oddHeader.left.text = "Quasar"
    ws.oddFooter.center.text = "Liquidación de asistencia  ·  página &P de &N"


def build_liquidacion_xlsx(
    db: Session,
    *,
    date_from: date,
    date_to: date,
    obra_id: Optional[int] = None,
) -> bytes:
    if date_to < date_from:
        raise ValueError("La fecha final no puede ser anterior a la inicial.")
    if (date_to - date_from).days > 400:
        raise ValueError("El periodo no puede superar 400 días.")

    obra_name = "Todas las obras"
    if obra_id is not None:
        obra = db.query(Obra).filter(Obra.id == obra_id).first()
        if not obra:
            raise ValueError(f"No existe la obra {obra_id}.")
        obra_name = obra.name

    rows = (
        db.query(ObraDailyAttendance)
        .options(
            joinedload(ObraDailyAttendance.obra),
            joinedload(ObraDailyAttendance.worker).joinedload(Worker.person),
            joinedload(ObraDailyAttendance.worker).joinedload(Worker.worker_position),
        )
        .filter(
            ObraDailyAttendance.removed_at.is_(None),
            ObraDailyAttendance.work_date >= date_from,
            ObraDailyAttendance.work_date <= date_to,
        )
        .order_by(
            ObraDailyAttendance.work_date.asc(),
            ObraDailyAttendance.id.asc(),
        )
    )
    if obra_id is not None:
        rows = rows.filter(ObraDailyAttendance.obra_id == obra_id)
    attendance = rows.all()

    by_worker: Dict[int, dict] = {}
    order: List[int] = []
    for row in attendance:
        worker = row.worker
        bucket = by_worker.get(row.worker_id)
        if bucket is None:
            person = worker.person if worker else None
            rate = worker.daily_rate if worker and worker.daily_rate is not None else None
            bucket = {
                "document": person.document_number if person else "",
                "name": person.name if person else f"Trabajador #{row.worker_id}",
                "position": (
                    worker.worker_position.name
                    if worker and worker.worker_position
                    else ""
                ),
                "account": worker.bank_account if worker and worker.bank_account else "",
                "rate": float(rate) if rate is not None else None,
                "dates": set(),
                "obras": defaultdict(int),
            }
            by_worker[row.worker_id] = bucket
            order.append(row.worker_id)
        bucket["dates"].add(row.work_date)
        site = row.obra.name if row.obra else f"Obra {row.obra_id}"
        bucket["obras"][site] += 1

    order.sort(key=lambda worker_id: by_worker[worker_id]["name"].lower())
    subtitle = f"{_fecha(date_from)}  —  {_fecha(date_to)}   ·   {obra_name}"

    book = Workbook()
    summary = book.active
    summary.title = "Liquidación"
    detail = book.create_sheet("Día a día")

    _header_band(summary, 8, "Liquidación de asistencia", subtitle)
    note = summary.cell(
        6,
        1,
        "El total es el valor día del trabajador × los días distintos con asistencia. "
        "Si un día estuvo en más de una obra, cuenta como un día y el detalle queda en «Día a día».",
    )
    summary.merge_cells(start_row=6, start_column=1, end_row=6, end_column=8)
    note.font = Font(name="Calibri", size=10, italic=True, color=MUTED)
    note.alignment = Alignment(wrap_text=True, vertical="center")
    summary.row_dimensions[6].height = 32

    labels = [
        "Documento",
        "Trabajador",
        "Cargo",
        "Cuenta Nequi",
        "Valor día",
        "Días",
        "Dónde estuvo",
        "Total a pagar",
    ]
    _table_header(summary, 8, labels)
    total_days = 0
    total_money = 0.0
    missing_rate = 0
    for index, worker_id in enumerate(order):
        item = by_worker[worker_id]
        days = len(item["dates"])
        total_days += days
        where = ", ".join(
            f"{name} ({count})" if count else name
            for name, count in sorted(item["obras"].items())
        )
        if item["rate"] is None:
            payable = "Sin valor día"
            missing_rate += 1
        else:
            payable = round(item["rate"] * days)
            total_money += payable
        _write_row(
            summary,
            9 + index,
            [
                item["document"] or "—",
                item["name"],
                item["position"] or "—",
                item["account"] or "—",
                item["rate"] if item["rate"] is not None else "—",
                days,
                where,
                payable,
            ],
            zebra=index % 2 == 1,
            money_cols=(5, 8),
        )
        summary.cell(9 + index, 6).alignment = Alignment(horizontal="center", vertical="center")

    if not order:
        _write_row(
            summary,
            9,
            ["Sin asistencias en este periodo", "", "", "", "", "", "", ""],
            zebra=False,
        )
        summary.merge_cells(start_row=9, start_column=1, end_row=9, end_column=8)
        end_row = 9
    else:
        end_row = 9 + len(order)
        payable_total = (
            "—" if missing_rate == len(order) else round(total_money)
        )
        _totals_row(
            summary,
            end_row,
            8,
            "TOTAL LIQUIDACIÓN",
            total_days,
            payable_total,
        )
        if missing_rate:
            if missing_rate == 1:
                warning = (
                    "1 trabajador no tiene valor día en RRHH. "
                    "Su fila dice «Sin valor día» hasta que se registre."
                )
            else:
                warning = (
                    f"{missing_rate} trabajadores no tienen valor día en RRHH. "
                    "Su fila dice «Sin valor día» hasta que se registre."
                )
            warn = summary.cell(end_row + 2, 1, warning)
            summary.merge_cells(
                start_row=end_row + 2,
                start_column=1,
                end_row=end_row + 2,
                end_column=8,
            )
            warn.font = Font(name="Calibri", size=10, italic=True, color="9A3412")
            warn.alignment = Alignment(wrap_text=True, vertical="center")
            summary.row_dimensions[end_row + 2].height = 28

    _widths(summary, [16, 32, 18, 18, 16, 10, 42, 18])
    summary.auto_filter.ref = f"A8:H{max(8, end_row - (0 if order else 1))}"
    if order:
        summary.auto_filter.ref = f"A8:H{end_row - 1}"
    _sheet_setup(summary, landscape=True)
    summary.page_setup.orientation = "landscape"
    summary.print_title_rows = "1:8"
    summary.page_setup.paperSize = summary.PAPERSIZE_TABLOID

    _header_band(detail, 7, "Dónde estuvo cada día", subtitle)
    detail_labels = [
        "Fecha",
        "Obra",
        "Documento",
        "Trabajador",
        "Cargo",
        "Cuenta Nequi",
        "Valor día",
    ]
    _table_header(detail, 7, detail_labels)
    if not attendance:
        _write_row(detail, 8, ["Sin asistencias en este periodo", "", "", "", "", "", ""], zebra=False)
        detail.merge_cells(start_row=8, start_column=1, end_row=8, end_column=7)
    else:
        for index, row in enumerate(attendance):
            worker = row.worker
            person = worker.person if worker else None
            rate = (
                float(worker.daily_rate)
                if worker and worker.daily_rate is not None
                else "—"
            )
            _write_row(
                detail,
                8 + index,
                [
                    row.work_date,
                    row.obra.name if row.obra else f"Obra {row.obra_id}",
                    person.document_number if person and person.document_number else "—",
                    person.name if person else f"Trabajador #{row.worker_id}",
                    worker.worker_position.name if worker and worker.worker_position else "—",
                    worker.bank_account if worker and worker.bank_account else "—",
                    rate,
                ],
                zebra=index % 2 == 1,
                money_cols=(7,),
            )
            detail.cell(8 + index, 1).number_format = "DD/MM/YYYY"
        detail.auto_filter.ref = f"A7:G{7 + len(attendance)}"
    _widths(detail, [14, 36, 16, 32, 18, 18, 16])
    _sheet_setup(detail, landscape=True)
    detail.print_title_rows = "1:7"

    book.properties.title = "Liquidación de asistencia"
    book.properties.creator = "Quasar"
    summary.sheet_properties.tabColor = BLUE
    detail.sheet_properties.tabColor = "1B2838"

    output = BytesIO()
    book.save(output)
    return output.getvalue()
